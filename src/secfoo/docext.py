"""Best-effort plain-text extraction from uploaded vendor documents (PDF/
DOCX/PPTX) for the Third-Party Risk Assessment skill. Never raises on a
malformed or unsupported file -- extraction failure degrades to an empty
string so an upload always succeeds even when its content can't be
analyzed; the original file is still stored as reference material either
way (see web/routes/assessments.py).

Deliberately pure-Python libraries only (pypdf, python-docx, python-pptx)
-- no system binaries (no poppler, no LibreOffice) so the Docker/
Kubernetes image needs no extra OS packages.

Known residual risk, accepted for now rather than silently claimed as
solved: these libraries fully load a file's internal structure into
memory (pypdf's page objects; python-docx/python-pptx's underlying
zipfile-based XML parsing). MAX_UPLOAD_BYTES bounds the raw file size, but
a small, deliberately crafted file with a high internal
compression/expansion ratio (a "zip bomb"-shaped .docx/.pptx, or a
pathological PDF object graph) could still cost more memory/CPU than its
file size implies. A stronger mitigation (subprocess isolation with a
hard resource/time limit) is a reasonable follow-up hardening item, not
built here.
"""

from __future__ import annotations

from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx"}

# Generous for a real vendor security document (SOC 2 reports, ISO
# certificates, and pen test reports are almost always well under this),
# while bounding the worst case of the residual risk noted above.
MAX_UPLOAD_BYTES = 25 * 1024 * 1024


def is_extractable(filename: str) -> bool:
    return Path(filename).suffix.lower() in SUPPORTED_EXTENSIONS


def extract_text(path: Path) -> str:
    """Extracts plain text from a PDF/DOCX/PPTX file at `path`. Returns
    "" for an unsupported extension, a file that fails to parse, or one
    with no extractable text (e.g. a scanned/image-only PDF -- this
    module does no OCR) -- callers treat "" as "nothing to analyze from
    this document," not as an error.
    """
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            return _extract_pdf(path)
        if suffix == ".docx":
            return _extract_docx(path)
        if suffix == ".pptx":
            return _extract_pptx(path)
    except Exception:
        return ""
    return ""


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text)
    return "\n\n".join(pages)


def _extract_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                slide_parts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        slide_parts.append(" | ".join(cells))
        if slide_parts:
            parts.append(f"--- Slide {slide_number} ---\n" + "\n".join(slide_parts))
    return "\n\n".join(parts)
